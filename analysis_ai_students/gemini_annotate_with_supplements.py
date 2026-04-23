"""
Annotate 10 manual-review papers using Gemini 3.0 Flash Preview.
Feeds BOTH the main paper text AND extracted supplement content into the model.

Usage:
    source ../.venv/bin/activate
    python gemini_annotate_with_supplements.py
"""

import os
import sys
import json
import time
import re
import zipfile
import xml.etree.ElementTree as ET
from string import Template
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
PROCESSED_PAPERS = BASE_DIR / "processed_papers.json"
PAPERS_DIR = BASE_DIR / "downloaded_papers"
SUPPL_DIR = BASE_DIR / "downloaded_supplements"
OUTPUT_DIR = Path(__file__).resolve().parent / "gemini_flash_results"

GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "")

TARGET_PAPERS = {
    "GSE234729": {"pmid": "37949031",  "pmcid": "PMC10843761"},
    "GSE98224":  {"pmid": "29507646",  "pmcid": "PMC5833042"},
    "GSE220877": {"pmid": "36744021",  "pmcid": "PMC9896899"},
    "GSE145357": {"pmid": "32640239",  "pmcid": "PMC7396155"},
    "GSE155750": {"pmid": "33784241",  "pmcid": "PMC8403268"},
    "GSE130339": {"pmid": "31658584",  "pmcid": "PMC6829352"},
    "GSE215421": {"pmid": "30611556",  "pmcid": "PMC7156023"},
    "GSE131729": {"pmid": "31917811",  "pmcid": "PMC6952106"},
    "GSE154829": {"pmid": "32721520",  "pmcid": "PMC7855285"},
    "GSE128381": {"pmid": "31110514",  "pmcid": "PMC6501552"},
}

EXTRACTION_PROMPT = Template('''\
You are an expert biomedical data extractor specializing in placental and pregnancy research.

Your task: read the paper text AND any supplementary material below, then answer each question.

IMPORTANT: Supplementary tables and methods often contain critical metadata
(demographics, gestational ages, birthweights, delivery modes, etc.) that may NOT
appear in the main paper text. You MUST check the supplementary material carefully.

Return ONLY a single valid JSON object. No markdown fences, no commentary.

=== OUTPUT FORMAT ===

For EACH question, return an object with these fields:
  "answer":     your answer (see rules below)
  "confidence": "high", "medium", or "low"
  "reason":     a 1-2 sentence explanation of WHY you chose this answer
  "evidence":   an array of supporting quotes from the paper or supplements

Example:
{
  "Birthweight of offspring provided (yes/no)": {
    "answer": "Yes",
    "confidence": "high",
    "reason": "Supplemental Table 1 reports birthweight for each sample in the cohort.",
    "evidence": [
      {"quote": "Mean birthweight was 3245 ± 412 g in the control group", "source": "Supplemental Table 1"},
      {"quote": "Birth weight (g) ... 2890 ± 680", "source": "Table 2, main text"}
    ]
  },
  "Paternal Weight provided (yes/no)": {
    "answer": "No",
    "confidence": "high",
    "reason": "Neither the paper nor supplements mention paternal weight.",
    "evidence": []
  }
}

=== ANSWER RULES ===

- Yes/No questions: return exactly "Yes" or "No".
  "Yes" means the paper OR its supplements EXPLICITLY provides or reports this data.
  "No" means it is not reported anywhere. If unclear or ambiguous, return "No".

- Pregnancy trimester: return one of "1st", "2nd", "3rd", "Term", "Premature",
  or an array if multiple are studied (e.g., ["1st", "3rd"]).
  "Term" = full-term delivery (~37-42 weeks).
  "Premature" = delivery before 37 weeks due to complications.

- GA fields (weeks): return a number or "Not Provided".

- List fields: return a JSON array of strings. Empty array [] if none.

- Free text fields: return the most specific value. "Not Provided" if absent.

- Do NOT invent data. Only report what is explicitly stated.

=== EVIDENCE RULES ===

- "quote": copy the EXACT text from the paper or supplement (keep it short, 1-2 sentences max)
- "source": be as specific as possible about where the quote comes from. Examples:
    "Main text, Methods section"
    "Main text, Results, Table 2"
    "Main text, Abstract"
    "Supplemental Table 1 (NIHMS1945468-supplement-Supplemental_Table.docx)"
    "Supplemental Methods"
  Include the section name, table/figure number, and supplement filename when applicable.
- For "No" answers: evidence should be an empty array []
- Include 1-3 evidence quotes per answer (the strongest supporting quotes)

=== CONFIDENCE GUIDELINES ===

- "high":   clearly and explicitly stated in paper or supplements
- "medium": requires some inference but well-supported
- "low":    ambiguous, contradictory, or incomplete information

=== QUESTIONS ===

1.  "Supervisor/Contact/Corresponding author name"
2.  "Supervisor/Contact/Corresponding author email"
3.  "Main topic of the publication"
4.  "Pregnancy trimester (1st, 2nd, 3rd, term, premature)"
5.  "Birthweight of offspring provided (yes/no)"
6.  "Gestational Age at delivery provided (yes/no)"
7.  "GA at delivery (weeks)"
8.  "Gestational Age at sample collection provided (yes/no)"
9.  "GA at sample collection (weeks)"
10. "Sex of Offspring Provided (yes/no)"
11. "Parity provided (yes/no)"
12. "Gravidity provided (yes/no)"
13. "Number of offspring per pregnancy provided (yes/no)"
14. "Self-reported race/ethnicity of mother provided (yes/no)"
15. "Genetic ancestry or genetic strain provided (yes/no)"
16. "Maternal Height provided (yes/no)"
17. "Maternal Pre-pregnancy Weight provided (yes/no)"
18. "Paternal Height provided (yes/no)"
19. "Paternal Weight provided (yes/no)"
20. "Maternal age at sample collection provided (yes/no)"
21. "Paternal age at sample collection provided (yes/no)"
22. "Samples from pregnancy complications collected"
23. "Mode of delivery provided (yes/no)"
24. "Pregnancy complications in data set (list)"
25. "Fetal complications listed (yes/no)"
26. "Fetal complications in data set (list)"
27. "Other Phenotypes Provided (list)"
28. "Hospital/Center where samples were collected"
29. "Country where samples were collected"

=== MAIN PAPER TEXT ===

$paper_text

=== SUPPLEMENTARY MATERIAL ===

$supplement_text
''')



def extract_docx_text(path: str) -> str:
    """Extract text from a .docx file (it's a zip of XML)."""
    try:
        with zipfile.ZipFile(path) as z:
            with z.open("word/document.xml") as f:
                tree = ET.parse(f)
                root = tree.getroot()
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                paragraphs = root.findall(".//w:p", ns)
                lines = []
                for p in paragraphs:
                    texts = [t.text for t in p.findall(".//w:t", ns) if t.text]
                    line = "".join(texts).strip()
                    if line:
                        lines.append(line)

                # also extract tables
                tables = root.findall(".//w:tbl", ns)
                for t_idx, table in enumerate(tables):
                    lines.append(f"\n[Table {t_idx + 1}]")
                    rows = table.findall(".//w:tr", ns)
                    for row in rows:
                        cells = row.findall(".//w:tc", ns)
                        cell_texts = []
                        for cell in cells:
                            ct = "".join(
                                t.text for t in cell.findall(".//w:t", ns) if t.text
                            )
                            cell_texts.append(ct.strip())
                        lines.append(" | ".join(cell_texts))

                return "\n".join(lines)
    except Exception as e:
        return f"[Error reading {os.path.basename(path)}: {e}]"


def extract_xlsx_text(path: str, max_rows: int = 200) -> str:
    """Extract text from .xlsx — reads sheet names + first N rows of each sheet."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"\n[Sheet: {sheet_name}]")
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                if row_count >= max_rows:
                    parts.append(f"  ... (truncated at {max_rows} rows)")
                    break
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    parts.append(" | ".join(cells))
                row_count += 1
        wb.close()
        return "\n".join(parts)
    except Exception as e:
        return f"[Error reading {os.path.basename(path)}: {e}]"


def extract_pdf_text(path: str) -> str:
    try:
        import fitz
        doc = fitz.open(path)
        parts = []
        for page_num, page in enumerate(doc, 1):
            text = page.get_text()
            if text.strip():
                parts.append(f"[Page {page_num}]\n{text}")
        doc.close()
        return "\n\n".join(parts)
    except Exception as e:
        return f"[Error reading {os.path.basename(path)}: {e}]"


def load_supplements(pmcid: str) -> str:
    """Load and concatenate all supplement text for a paper."""
    suppl_path = SUPPL_DIR / pmcid
    if not suppl_path.exists():
        return "(No supplementary files available)"

    parts = []
    for fname in sorted(os.listdir(suppl_path)):
        fpath = str(suppl_path / fname)
        ext = os.path.splitext(fname)[1].lower()

        parts.append(f"\n--- Supplement file: {fname} ---")
        if ext in (".docx", ".doc"):
            parts.append(extract_docx_text(fpath))
        elif ext in (".xlsx", ".xls"):
            parts.append(extract_xlsx_text(fpath))
        elif ext == ".pdf":
            parts.append(extract_pdf_text(fpath))
        elif ext in (".csv", ".txt"):
            try:
                with open(fpath, encoding="utf-8", errors="replace") as f:
                    parts.append(f.read()[:50000])
            except Exception as e:
                parts.append(f"[Error reading {fname}: {e}]")
        else:
            parts.append(f"[Skipped: unsupported format {ext}]")

    text = "\n".join(parts)
    # truncate if too long (gemini context limit)
    if len(text) > 80000:
        text = text[:80000] + "\n\n... (supplementary text truncated at 80,000 chars)"
    return text



MAX_RETRIES = 5
RETRY_DELAY = 2.0


def call_gemini(prompt: str) -> str:
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=GEMINI_API_KEY)
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = client.models.generate_content(
                model="gemini-3-flash-preview",
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json"
                ),
            )
            return resp.text
        except Exception as e:
            print(f"    Attempt {attempt}/{MAX_RETRIES}: {e}")
            if attempt < MAX_RETRIES:
                time.sleep(RETRY_DELAY * attempt)
    raise RuntimeError(f"Gemini failed after {MAX_RETRIES} retries")


def parse_json_response(text: str) -> dict:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    return json.loads(text)



def main():
    if not GEMINI_API_KEY:
        print("ERROR: Set GEMINI_API_KEY environment variable")
        print("  export GEMINI_API_KEY='your-key-here'")
        sys.exit(1)

    # load processed paper chunks
    with open(PROCESSED_PAPERS, encoding="utf-8") as f:
        all_papers = json.load(f)
    paper_lookup = {}
    for p in all_papers:
        paper_lookup[p["pmcid"]] = " ".join(p.get("chunks", []))

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    results = {}
    for geo_id, info in TARGET_PAPERS.items():
        pmcid = info["pmcid"]
        print(f"\n{'='*60}")
        print(f"Processing {geo_id} ({pmcid})")
        print(f"{'='*60}")

        # check if already done
        out_file = OUTPUT_DIR / f"{geo_id}.json"
        if out_file.exists():
            print("  Already processed, loading existing results...")
            with open(out_file) as f:
                data = json.load(f)
            # unwrap if gemini returned [{...}] instead of {...}
            if isinstance(data, list) and len(data) == 1 and isinstance(data[0], dict):
                data = data[0]
            results[geo_id] = data
            continue

        # get main paper text
        paper_text = paper_lookup.get(pmcid, "")
        if not paper_text:
            print(f"  WARNING: No paper text found for {pmcid} in processed_papers.json")
            continue

        # get supplement text
        print("  Loading supplements...")
        supplement_text = load_supplements(pmcid)
        suppl_len = len(supplement_text)
        print(f"  Supplement text: {suppl_len:,} chars")

        # build prompt
        prompt = EXTRACTION_PROMPT.substitute(
            paper_text=paper_text[:100000],  # cap main text
            supplement_text=supplement_text,
        )
        print(f"  Total prompt: {len(prompt):,} chars")

        # call gemini
        print("  Calling Gemini 3.0 Flash Preview...")
        raw_response = call_gemini(prompt)

        # save raw response
        raw_file = OUTPUT_DIR / f"{geo_id}.raw.txt"
        with open(raw_file, "w") as f:
            f.write(raw_response)

        # parse
        try:
            parsed = parse_json_response(raw_response)
            if isinstance(parsed, list) and len(parsed) == 1 and isinstance(parsed[0], dict):
                parsed = parsed[0]
            results[geo_id] = parsed
            with open(out_file, "w") as f:
                json.dump(parsed, f, indent=2)
            print(f"  OK — {len(parsed)} fields extracted")
        except json.JSONDecodeError as e:
            print(f"  ERROR parsing JSON: {e}")
            print(f"  Raw response saved to {raw_file}")
            continue

        # rate limit
        time.sleep(1)

    # build summary excel
    print(f"\n{'='*60}")
    print("Building summary Excel...")
    print(f"{'='*60}")

    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gemini Flash Annotations"

    # headers matching manual review format
    QUESTIONS = [
        "Supervisor/Contact/Corresponding author name",
        "Supervisor/Contact/Corresponding author email",
        "Main topic of the publication",
        "Pregnancy trimester (1st, 2nd, 3rd, term, premature)",
        "Birthweight of offspring provided (yes/no)",
        "Gestational Age at delivery provided (yes/no)",
        "GA at delivery (weeks)",
        "Gestational Age at sample collection provided (yes/no)",
        "GA at sample collection (weeks)",
        "Sex of Offspring Provided (yes/no)",
        "Parity provided (yes/no)",
        "Gravidity provided (yes/no)",
        "Number of offspring per pregnancy provided (yes/no)",
        "Self-reported race/ethnicity of mother provided (yes/no)",
        "Genetic ancestry or genetic strain provided (yes/no)",
        "Maternal Height provided (yes/no)",
        "Maternal Pre-pregnancy Weight provided (yes/no)",
        "Paternal Height provided (yes/no)",
        "Paternal Weight provided (yes/no)",
        "Maternal age at sample collection provided (yes/no)",
        "Paternal age at sample collection provided (yes/no)",
        "Samples from pregnancy complications collected",
        "Mode of delivery provided (yes/no)",
        "Pregnancy complications in data set (list)",
        "Fetal complications listed (yes/no)",
        "Fetal complications in data set (list)",
        "Other Phenotypes Provided (list)",
        "Hospital/Center where samples were collected",
        "Country where samples were collected",
    ]

    headers = ["GEO ID", "PMCID", "PMID"] + QUESTIONS
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # data rows
    row_num = 2
    for geo_id, info in TARGET_PAPERS.items():
        ws.cell(row=row_num, column=1, value=geo_id)
        ws.cell(row=row_num, column=2, value=info["pmcid"])
        ws.cell(row=row_num, column=3, value=info["pmid"])

        parsed = results.get(geo_id, {})
        for q_idx, q in enumerate(QUESTIONS):
            val = parsed.get(q, {})
            if isinstance(val, dict):
                answer = val.get("answer", "")
            else:
                answer = val
            if isinstance(answer, list):
                answer = ", ".join(str(a) for a in answer)
            ws.cell(row=row_num, column=4 + q_idx, value=str(answer))

        row_num += 1

    # auto-width columns
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    # sheet 2: confidence + reason
    ws2 = wb.create_sheet("Confidence & Reasoning")
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font

    row_num = 2
    for geo_id, info in TARGET_PAPERS.items():
        ws2.cell(row=row_num, column=1, value=geo_id)
        ws2.cell(row=row_num, column=2, value=info["pmcid"])
        ws2.cell(row=row_num, column=3, value=info["pmid"])

        parsed = results.get(geo_id, {})
        for q_idx, q in enumerate(QUESTIONS):
            val = parsed.get(q, {})
            if isinstance(val, dict):
                conf = val.get("confidence", "")
                reason = val.get("reason", "")
                cell_val = f"[{conf}] {reason}" if reason else conf
            else:
                cell_val = ""
            cell = ws2.cell(row=row_num, column=4 + q_idx, value=cell_val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            # color code confidence
            if "high" in str(conf).lower():
                cell.fill = PatternFill(start_color="C6EFCE", fill_type="solid")
            elif "medium" in str(conf).lower():
                cell.fill = PatternFill(start_color="FFEB9C", fill_type="solid")
            elif "low" in str(conf).lower():
                cell.fill = PatternFill(start_color="FFC7CE", fill_type="solid")

        row_num += 1

    # sheet 3: evidence (quotes + sources)
    ws3 = wb.create_sheet("Evidence")
    ev_headers = ["GEO ID", "Question", "Answer", "Confidence", "Reason", "Quote", "Source"]
    for col, h in enumerate(ev_headers, 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font

    ev_row = 2
    for geo_id in TARGET_PAPERS:
        parsed = results.get(geo_id, {})
        for q in QUESTIONS:
            val = parsed.get(q, {})
            if not isinstance(val, dict):
                continue

            answer = val.get("answer", "")
            if isinstance(answer, list):
                answer = ", ".join(str(a) for a in answer)
            confidence = val.get("confidence", "")
            reason = val.get("reason", "")
            evidence = val.get("evidence", [])

            if not evidence:
                # still write a row with the answer/reason, just no quote
                ws3.cell(row=ev_row, column=1, value=geo_id)
                ws3.cell(row=ev_row, column=2, value=q)
                ws3.cell(row=ev_row, column=3, value=str(answer))
                ws3.cell(row=ev_row, column=4, value=confidence)
                ws3.cell(row=ev_row, column=5, value=reason)
                ws3.cell(row=ev_row, column=6, value="")
                ws3.cell(row=ev_row, column=7, value="")
                ev_row += 1
            else:
                for ev in evidence:
                    quote = ev.get("quote", "") if isinstance(ev, dict) else str(ev)
                    source = ev.get("source", "") if isinstance(ev, dict) else ""
                    ws3.cell(row=ev_row, column=1, value=geo_id)
                    ws3.cell(row=ev_row, column=2, value=q)
                    ws3.cell(row=ev_row, column=3, value=str(answer))
                    ws3.cell(row=ev_row, column=4, value=confidence)
                    ws3.cell(row=ev_row, column=5, value=reason)
                    cell_q = ws3.cell(row=ev_row, column=6, value=quote)
                    cell_q.alignment = Alignment(wrap_text=True)
                    ws3.cell(row=ev_row, column=7, value=source)
                    ev_row += 1

    # auto-width evidence sheet
    for col in ws3.columns:
        max_len = max(len(str(c.value or "")[:60]) for c in col)
        ws3.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    out_xlsx = OUTPUT_DIR / "gemini_flash_annotated.xlsx"
    wb.save(out_xlsx)
    print(f"\nSaved: {out_xlsx}")
    print(f"Papers processed: {len(results)}/{len(TARGET_PAPERS)}")


if __name__ == "__main__":
    main()
