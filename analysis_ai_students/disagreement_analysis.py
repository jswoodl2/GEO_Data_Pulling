"""
Build a disagreement table comparing Gemini Flash annotations vs student majority vote.
Shows every yes/no disagreement with Gemini's evidence quotes and sources.

Requires:
    - student_majority_vote.xlsx (run student_majority_vote.py first)
    - gemini_flash_results/*.json (run gemini_annotate_with_supplements.py first)

Usage:
    source ../.venv/bin/activate
    python disagreement_analysis.py
"""

import json
import os
from collections import Counter

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BASE_DIR = os.path.dirname(__file__)
STUDENT_FILE = os.path.join(BASE_DIR, "student_majority_vote.xlsx")
GEMINI_DIR = os.path.join(BASE_DIR, "gemini_flash_results")
OUTPUT_FILE = os.path.join(BASE_DIR, "disagreement_analysis.xlsx")

YESNO = [
    "Birthweight of offspring provided (yes/no)",
    "Gestational Age at delivery provided (yes/no)",
    "Gestational Age at sample collection provided (yes/no)",
    "Sex of Offspring Provided (yes/no)",
    "Parity provided (yes/no)",
    "Gravidity provided (yes/no)",
    "Number of offspring per pregnancy provided (yes/no)",
    "Self-reported race/ethnicity of mother provided (yes/no)",
    "Genetic ancestry or genetic strain provided (yes/no)",
    "Maternal Height provided (yes/no)",
    "Maternal Pre-pregnancy Weight provided (yes/no)",
    "Paternal Height provided (yes/no)",
    "Paternal Weight provided (yes/no)",
    "Maternal age at sample collection provided (yes/no)",
    "Paternal age at sample collection provided (yes/no)",
    "Samples from pregnancy complications collected",
    "Mode of delivery provided (yes/no)",
    "Fetal complications listed (yes/no)",
]


def norm(v):
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in ("yes", "y", "true", "1"):
        return "Yes"
    if s in ("no", "n", "false", "0"):
        return "No"
    return None


def load_student_data():
    wb = openpyxl.load_workbook(STUDENT_FILE)

    # majority vote
    ws = wb["Majority Vote"]
    headers = [c.value for c in ws[1]]
    student_data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        geo_id = row[0]
        if not geo_id:
            continue
        student_data[geo_id] = dict(zip(headers, row))

    # individual votes
    ws_d = wb["Vote Details"]
    vote_detail = {}
    for row in ws_d.iter_rows(min_row=2, values_only=True):
        geo, q, s1, s2, s3, consensus, agreement = row[:7]
        if geo and q:
            vote_detail[(geo, q)] = {"Set1": s1, "Set2": s2, "Set3": s3}

    return student_data, vote_detail


def load_gemini_data():
    gemini_data = {}
    for f in os.listdir(GEMINI_DIR):
        if f.endswith(".json"):
            geo_id = f.replace(".json", "")
            with open(os.path.join(GEMINI_DIR, f)) as fh:
                data = json.load(fh)
            if isinstance(data, list) and len(data) == 1:
                data = data[0]
            gemini_data[geo_id] = data
    return gemini_data


def main():
    student_data, vote_detail = load_student_data()
    gemini_data = load_gemini_data()

    # styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    gemini_yes_fill = PatternFill(start_color="FFC7CE", fill_type="solid")
    student_yes_fill = PatternFill(start_color="FFEB9C", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(bottom=Side(style="thin", color="CCCCCC"))

    wb = openpyxl.Workbook()

    # sheet 1: disagreement table with evidence
    ws1 = wb.active
    ws1.title = "Disagreements + Evidence"
    headers = [
        "GEO ID", "Question",
        "Student\n(majority)", "Set1", "Set2", "Set3",
        "Gemini", "Confidence",
        "Gemini Reason",
        "Evidence Quote(s)", "Evidence Source(s)",
    ]

    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap

    row_num = 2
    disagreements = []
    total_comparisons = 0
    total_agree = 0

    for geo_id in sorted(student_data.keys()):
        g = gemini_data.get(geo_id, {})
        s = student_data[geo_id]

        for q in YESNO:
            s_val = norm(s.get(q))
            g_entry = g.get(q, {})
            if isinstance(g_entry, dict):
                g_val = norm(g_entry.get("answer"))
                confidence = g_entry.get("confidence", "")
                reason = g_entry.get("reason", "")
                evidence = g_entry.get("evidence", [])
            else:
                g_val = norm(g_entry)
                confidence = reason = ""
                evidence = []

            if s_val is None or g_val is None:
                continue

            total_comparisons += 1
            if s_val == g_val:
                total_agree += 1
                continue

            # get individual student votes
            vd = vote_detail.get((geo_id, q), {})
            s1 = vd.get("Set1", "")
            s2 = vd.get("Set2", "")
            s3 = vd.get("Set3", "")

            # format evidence
            quotes = []
            sources = []
            for ev in evidence or []:
                if isinstance(ev, dict):
                    quotes.append(ev.get("quote", ""))
                    sources.append(ev.get("source", ""))

            quotes_str = "\n---\n".join(quotes) if quotes else "(no quotes)"
            sources_str = "\n".join(sources) if sources else ""

            direction = "Gemini=Yes" if g_val == "Yes" else "Student=Yes"

            ws1.cell(row=row_num, column=1, value=geo_id)
            ws1.cell(row=row_num, column=2, value=q).alignment = wrap
            cell_s = ws1.cell(row=row_num, column=3, value=s_val)
            ws1.cell(row=row_num, column=4, value=s1)
            ws1.cell(row=row_num, column=5, value=s2)
            ws1.cell(row=row_num, column=6, value=s3)
            cell_g = ws1.cell(row=row_num, column=7, value=g_val)
            ws1.cell(row=row_num, column=8, value=confidence)
            ws1.cell(row=row_num, column=9, value=reason).alignment = wrap
            ws1.cell(row=row_num, column=10, value=quotes_str).alignment = wrap
            ws1.cell(row=row_num, column=11, value=sources_str).alignment = wrap

            if g_val == "Yes":
                cell_g.fill = gemini_yes_fill
            else:
                cell_s.fill = student_yes_fill

            for c in range(1, 12):
                ws1.cell(row=row_num, column=c).border = thin_border

            disagreements.append({
                "geo": geo_id, "q": q, "student": s_val, "gemini": g_val,
                "direction": direction, "confidence": confidence,
            })
            row_num += 1

    # sheet 2: summary
    ws2 = wb.create_sheet("Summary")

    gemini_yes = sum(1 for d in disagreements if d["direction"] == "Gemini=Yes")
    student_yes = sum(1 for d in disagreements if d["direction"] == "Student=Yes")
    q_counts = Counter(d["q"] for d in disagreements)
    paper_counts = Counter(d["geo"] for d in disagreements)

    ws2.cell(row=1, column=1, value="Metric").font = Font(bold=True)
    ws2.cell(row=1, column=2, value="Value").font = Font(bold=True)

    stats = [
        ("Total yes/no comparisons", total_comparisons),
        ("Agreements", total_agree),
        ("Agreement rate", f"{100 * total_agree / total_comparisons:.1f}%" if total_comparisons else "N/A"),
        ("Disagreements", len(disagreements)),
        ("", ""),
        ("Gemini=Yes, Student=No", f"{gemini_yes} ({100 * gemini_yes / len(disagreements):.0f}% of disagreements)" if disagreements else 0),
        ("Student=Yes, Gemini=No", f"{student_yes} ({100 * student_yes / len(disagreements):.0f}% of disagreements)" if disagreements else 0),
        ("", ""),
        ("DISAGREEMENTS BY QUESTION", ""),
    ]
    for q, cnt in q_counts.most_common():
        stats.append((q, cnt))

    stats.append(("", ""))
    stats.append(("DISAGREEMENTS BY PAPER", ""))
    for geo, cnt in paper_counts.most_common():
        stats.append((geo, cnt))

    for i, (metric, val) in enumerate(stats, 2):
        ws2.cell(row=i, column=1, value=metric)
        ws2.cell(row=i, column=2, value=val)
        if metric and metric.isupper():
            ws2.cell(row=i, column=1).font = Font(bold=True)

    # column widths
    ws1.column_dimensions["A"].width = 12
    ws1.column_dimensions["B"].width = 40
    ws1.column_dimensions["C"].width = 10
    ws1.column_dimensions["D"].width = 8
    ws1.column_dimensions["E"].width = 8
    ws1.column_dimensions["F"].width = 8
    ws1.column_dimensions["G"].width = 10
    ws1.column_dimensions["H"].width = 12
    ws1.column_dimensions["I"].width = 45
    ws1.column_dimensions["J"].width = 55
    ws1.column_dimensions["K"].width = 35
    ws2.column_dimensions["A"].width = 55
    ws2.column_dimensions["B"].width = 20

    wb.save(OUTPUT_FILE)
    print(f"Saved: {OUTPUT_FILE}")
    print(f"\nOverall: {total_agree}/{total_comparisons} agree ({100 * total_agree / total_comparisons:.1f}%)")
    print(f"Disagreements: {len(disagreements)}")
    print(f"  Gemini=Yes, Student=No: {gemini_yes}")
    print(f"  Student=Yes, Gemini=No: {student_yes}")
    print(f"\nTop disagreement questions:")
    for q, cnt in q_counts.most_common(5):
        print(f"  {cnt}x  {q}")
    print(f"\nTop disagreement papers:")
    for geo, cnt in paper_counts.most_common(5):
        print(f"  {cnt}x  {geo}")


if __name__ == "__main__":
    main()
