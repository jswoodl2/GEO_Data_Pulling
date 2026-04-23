"""
Re-query Gemini 3 Flash on just the disagreement questions, then build
an investigation sheet with GEO ID, Question, Answer, Reason, Quote, Source.

Reads: disagreement_analysis.xlsx (the 43 disagreements)
Writes: disagreement_investigation.xlsx

Usage:
    source ../.venv/bin/activate
    export GEMINI_API_KEY='your-key'
    python disagreement_investigation.py
"""

import os
import sys
import json
import time
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from string import Template
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

INVESTIGATION_DIR = Path(__file__).resolve().parent
ANALYSIS_DIR = INVESTIGATION_DIR.parent
BASE_DIR = ANALYSIS_DIR.parent
PROCESSED_PAPERS = BASE_DIR / "processed_papers.json"
SUPPL_DIR = BASE_DIR / "downloaded_supplements"
DISAGREEMENT_FILE = ANALYSIS_DIR / "disagreement_analysis.xlsx"
OUTPUT_DIR = INVESTIGATION_DIR / "gemini_recheck_json"
OUTPUT_XLSX = INVESTIGATION_DIR / "disagreement_investigation.xlsx"

GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "")

TARGET_PAPERS = {
    "GSE234729": {"pmid": "37949031", "pmcid": "PMC10843761"},
    "GSE98224":  {"pmid": "29507646", "pmcid": "PMC5833042"},
    "GSE220877": {"pmid": "36744021", "pmcid": "PMC9896899"},
    "GSE145357": {"pmid": "32640239", "pmcid": "PMC7396155"},
    "GSE155750": {"pmid": "33784241", "pmcid": "PMC8403268"},
    "GSE130339": {"pmid": "31658584", "pmcid": "PMC6829352"},
    "GSE215421": {"pmid": "30611556", "pmcid": "PMC7156023"},
    "GSE131729": {"pmid": "31917811", "pmcid": "PMC6952106"},
    "GSE154829": {"pmid": "32721520", "pmcid": "PMC7855285"},
    "GSE128381": {"pmid": "31110514", "pmcid": "PMC6501552"},
}


FOCUSED_PROMPT = Template('''\
You are an expert biomedical data extractor specializing in placental and pregnancy research.

You previously answered a set of yes/no questions for this paper. These specific
questions had disagreements with human annotators and I want a careful, focused
re-examination.

Read the paper text AND its supplementary material carefully. For each question,
decide your answer based ONLY on what is explicitly stated.

IMPORTANT: Supplementary tables and methods often contain critical metadata
(demographics, gestational ages, birthweights, delivery modes, etc.) that may NOT
appear in the main paper text. You MUST check the supplementary material carefully.

Return ONLY a single valid JSON object. No markdown fences, no commentary.

=== OUTPUT FORMAT ===

For EACH question below, return an object:
{
  "<question text>": {
    "answer": "Yes" or "No",
    "confidence": "high" | "medium" | "low",
    "reason": "1-2 sentence explanation of WHY you chose this answer",
    "evidence": [
      {"quote": "exact text from paper/supplement", "source": "where the quote is from"}
    ]
  },
  ...
}

=== ANSWER RULES ===

- "Yes" means the paper OR its supplements EXPLICITLY provides or reports this data.
- "No" means it is not reported anywhere. If unclear or ambiguous, return "No".
- Do NOT invent data. Only report what is explicitly stated.

=== EVIDENCE RULES ===

- "quote": copy the EXACT text from the paper or supplement (1-2 sentences max).
- "source": be specific: section name, table/figure number, supplement filename, etc.
  Examples: "Main text, Methods section", "Supplemental Table 1 (NIHMS1945468-supplement-Supplemental_Table.docx)"
- For "No" answers: evidence should be an empty array [].
- Include 1-3 strongest quotes per answer.

=== QUESTIONS TO RE-EXAMINE ===

$questions_list

=== MAIN PAPER TEXT ===

$paper_text

=== SUPPLEMENTARY MATERIAL ===

$supplement_text
''')



def extract_docx_text(path: str) -> str:
    try:
        with zipfile.ZipFile(path) as z:
            with z.open("word/document.xml") as f:
                tree = ET.parse(f)
                root = tree.getroot()
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                lines = []
                for p in root.findall(".//w:p", ns):
                    texts = [t.text for t in p.findall(".//w:t", ns) if t.text]
                    line = "".join(texts).strip()
                    if line:
                        lines.append(line)
                for t_idx, table in enumerate(root.findall(".//w:tbl", ns)):
                    lines.append(f"\n[Table {t_idx + 1}]")
                    for row in table.findall(".//w:tr", ns):
                        cell_texts = []
                        for cell in row.findall(".//w:tc", ns):
                            ct = "".join(t.text for t in cell.findall(".//w:t", ns) if t.text)
                            cell_texts.append(ct.strip())
                        lines.append(" | ".join(cell_texts))
                return "\n".join(lines)
    except Exception as e:
        return f"[Error reading {os.path.basename(path)}: {e}]"


def extract_xlsx_text(path: str, max_rows: int = 200) -> str:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"\n[Sheet: {sheet_name}]")
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                if row_count >= max_rows:
                    parts.append(f"  ... (truncated at {max_rows} rows)")
                    break
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    parts.append(" | ".join(cells))
                row_count += 1
        wb.close()
        return "\n".join(parts)
    except Exception as e:
        return f"[Error reading {os.path.basename(path)}: {e}]"


def extract_pdf_text(path: str) -> str:
    try:
        import fitz
        doc = fitz.open(path)
        parts = []
        for page_num, page in enumerate(doc, 1):
            text = page.get_text()
            if text.strip():
                parts.append(f"[Page {page_num}]\n{text}")
        doc.close()
        return "\n\n".join(parts)
    except Exception as e:
        return f"[Error reading {os.path.basename(path)}: {e}]"


def load_supplements(pmcid: str) -> str:
    suppl_path = SUPPL_DIR / pmcid
    if not suppl_path.exists():
        return "(No supplementary files available)"

    parts = []
    for fname in sorted(os.listdir(suppl_path)):
        fpath = str(suppl_path / fname)
        ext = os.path.splitext(fname)[1].lower()
        parts.append(f"\n--- Supplement file: {fname} ---")
        if ext in (".docx", ".doc"):
            parts.append(extract_docx_text(fpath))
        elif ext in (".xlsx", ".xls"):
            parts.append(extract_xlsx_text(fpath))
        elif ext == ".pdf":
            parts.append(extract_pdf_text(fpath))
        elif ext in (".csv", ".txt"):
            try:
                with open(fpath, encoding="utf-8", errors="replace") as f:
                    parts.append(f.read()[:50000])
            except Exception as e:
                parts.append(f"[Error: {e}]")

    text = "\n".join(parts)
    if len(text) > 80000:
        text = text[:80000] + "\n\n... (truncated at 80,000 chars)"
    return text



MAX_RETRIES = 5
RETRY_DELAY = 2.0


def call_gemini(prompt: str) -> str:
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=GEMINI_API_KEY)
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = client.models.generate_content(
                model="gemini-3-flash-preview",
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json"
                ),
            )
            return resp.text
        except Exception as e:
            print(f"    Attempt {attempt}/{MAX_RETRIES}: {e}")
            if attempt < MAX_RETRIES:
                time.sleep(RETRY_DELAY * attempt)
    raise RuntimeError(f"Gemini failed after {MAX_RETRIES} retries")


def parse_json_response(text: str) -> dict:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        decoder = json.JSONDecoder()
        data, _ = decoder.raw_decode(text)
    if isinstance(data, list) and len(data) == 1 and isinstance(data[0], dict):
        data = data[0]
    return data



def load_disagreements():
    wb = openpyxl.load_workbook(DISAGREEMENT_FILE)
    ws = wb["Disagreements + Evidence"]

    disagreements = defaultdict(list)  # geo_id -> [question, ...]
    for row in ws.iter_rows(min_row=2, values_only=True):
        geo_id, question = row[0], row[1]
        if geo_id and question:
            disagreements[geo_id].append(question)
    return disagreements



def main():
    if not GEMINI_API_KEY:
        print("ERROR: Set GEMINI_API_KEY environment variable")
        sys.exit(1)

    disagreements = load_disagreements()
    total_questions = sum(len(qs) for qs in disagreements.values())
    print(f"Papers with disagreements: {len(disagreements)}")
    print(f"Total disagreement questions: {total_questions}")

    # load paper text
    with open(PROCESSED_PAPERS, encoding="utf-8") as f:
        all_papers = json.load(f)
    paper_lookup = {p["pmcid"]: " ".join(p.get("chunks", [])) for p in all_papers}

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # run gemini per paper
    all_results = {}  # geo_id -> {question: {answer, reason, evidence}}
    for geo_id, questions in disagreements.items():
        pmcid = TARGET_PAPERS[geo_id]["pmcid"]
        print(f"\n{'=' * 60}")
        print(f"Re-querying {geo_id} ({pmcid}) — {len(questions)} questions")
        print(f"{'=' * 60}")

        # check cache
        cache_file = OUTPUT_DIR / f"{geo_id}_disagreement_recheck.json"
        if cache_file.exists():
            print("  Loading cached result...")
            with open(cache_file) as f:
                all_results[geo_id] = json.load(f)
            continue

        paper_text = paper_lookup.get(pmcid, "")
        if not paper_text:
            print(f"  WARNING: No paper text for {pmcid}")
            continue

        supplement_text = load_supplements(pmcid)
        questions_list = "\n".join(f"- \"{q}\"" for q in questions)

        prompt = FOCUSED_PROMPT.substitute(
            questions_list=questions_list,
            paper_text=paper_text[:100000],
            supplement_text=supplement_text,
        )
        print(f"  Prompt: {len(prompt):,} chars")

        print("  Calling Gemini 3.0 Flash...")
        raw = call_gemini(prompt)

        # save raw
        (OUTPUT_DIR / f"{geo_id}_disagreement_recheck.raw.txt").write_text(raw)

        try:
            parsed = parse_json_response(raw)
            all_results[geo_id] = parsed
            with open(cache_file, "w") as f:
                json.dump(parsed, f, indent=2)
            print(f"  OK — {len(parsed)} fields returned")
        except json.JSONDecodeError as e:
            print(f"  ERROR parsing JSON: {e}")
            continue

        time.sleep(1)

    # build excel
    print(f"\n{'=' * 60}")
    print("Building investigation Excel...")
    print(f"{'=' * 60}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Disagreement Investigation"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    wrap = Alignment(wrap_text=True, vertical="top")
    high_fill = PatternFill(start_color="C6EFCE", fill_type="solid")
    med_fill = PatternFill(start_color="FFEB9C", fill_type="solid")
    low_fill = PatternFill(start_color="FFC7CE", fill_type="solid")

    headers = [
        "GEO ID", "Question", "Gemini Answer", "Confidence",
        "Gemini Reason", "Evidence Quote", "Evidence Source",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap

    row_num = 2
    for geo_id in sorted(disagreements.keys()):
        questions = disagreements[geo_id]
        results = all_results.get(geo_id, {})

        for q in questions:
            entry = results.get(q, {})
            if isinstance(entry, dict):
                answer = entry.get("answer", "")
                confidence = entry.get("confidence", "")
                reason = entry.get("reason", "")
                evidence = entry.get("evidence", []) or []
            else:
                answer = str(entry)
                confidence = reason = ""
                evidence = []

            if isinstance(answer, list):
                answer = ", ".join(str(a) for a in answer)

            if not evidence:
                # one row with no evidence
                ws.cell(row=row_num, column=1, value=geo_id)
                ws.cell(row=row_num, column=2, value=q).alignment = wrap
                ws.cell(row=row_num, column=3, value=str(answer))
                conf_cell = ws.cell(row=row_num, column=4, value=confidence)
                ws.cell(row=row_num, column=5, value=reason).alignment = wrap
                ws.cell(row=row_num, column=6, value="")
                ws.cell(row=row_num, column=7, value="")
                if "high" in confidence.lower():
                    conf_cell.fill = high_fill
                elif "medium" in confidence.lower():
                    conf_cell.fill = med_fill
                elif "low" in confidence.lower():
                    conf_cell.fill = low_fill
                row_num += 1
            else:
                # one row per evidence item
                for ev in evidence:
                    quote = ev.get("quote", "") if isinstance(ev, dict) else str(ev)
                    source = ev.get("source", "") if isinstance(ev, dict) else ""
                    ws.cell(row=row_num, column=1, value=geo_id)
                    ws.cell(row=row_num, column=2, value=q).alignment = wrap
                    ws.cell(row=row_num, column=3, value=str(answer))
                    conf_cell = ws.cell(row=row_num, column=4, value=confidence)
                    ws.cell(row=row_num, column=5, value=reason).alignment = wrap
                    ws.cell(row=row_num, column=6, value=quote).alignment = wrap
                    ws.cell(row=row_num, column=7, value=source).alignment = wrap
                    if "high" in confidence.lower():
                        conf_cell.fill = high_fill
                    elif "medium" in confidence.lower():
                        conf_cell.fill = med_fill
                    elif "low" in confidence.lower():
                        conf_cell.fill = low_fill
                    row_num += 1

    # column widths
    widths = {"A": 12, "B": 40, "C": 12, "D": 12, "E": 45, "F": 55, "G": 35}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wb.save(OUTPUT_XLSX)
    print(f"\nSaved: {OUTPUT_XLSX}")
    print(f"Total rows: {row_num - 2}")


if __name__ == "__main__":
    main()
